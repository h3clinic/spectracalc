"""Berlman Spectra Digitizer — interactive web app.

Upload a fluorescence spectrum image, auto-digitize it with the Berlman
pipeline, then fine-tune by adding/erasing dots before exporting.  Every
edit regenerates, per curve: a CSV, an Excel workbook with a scatter chart,
an Excel-style chart image, and a PhotochemCAD-style formatted page.
"""
import io, os, sys, base64, json, csv, tempfile, uuid, re, shutil, time
import cv2
import numpy as np
from fastapi import FastAPI, UploadFile, File
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import berlman as B

OUTPUTS = os.path.join(os.path.dirname(__file__), "outputs")
CACHE = os.path.join(os.path.dirname(__file__), "cache")
os.makedirs(OUTPUTS, exist_ok=True)
os.makedirs(CACHE, exist_ok=True)

PAGES_DIR = os.path.join(os.path.dirname(__file__), "..", "berlman_run600", "pages")
GALLERY_DIR = next((d for d in (
    os.path.join(os.path.dirname(__file__), "..", "gallery"),     # local layout
    os.path.join(os.path.dirname(__file__), "..", "overlays"),    # repo layout
) if os.path.isdir(d)), "")
_HERE = os.path.dirname(__file__)
SPECTRA_JSON_CANDIDATES = [
    os.path.join(_HERE, "spectra_all.json"),              # bundled beside app
    os.path.join(_HERE, "..", "data", "spectra_all.json"),  # repo layout
    os.path.expanduser("~/Downloads/Berlman_600dpi_spectra/spectra_all.json"),
]

app = FastAPI(title="Berlman Spectra Digitizer")
app.mount("/static", StaticFiles(directory=os.path.join(os.path.dirname(__file__), "static")), name="static")
app.mount("/outputs", StaticFiles(directory=OUTPUTS), name="outputs")
if os.path.isdir(GALLERY_DIR):
    app.mount("/overlays", StaticFiles(directory=GALLERY_DIR), name="overlays")


@app.get("/", response_class=HTMLResponse)
async def index():
    with open(os.path.join(os.path.dirname(__file__), "static", "index.html")) as f:
        return f.read()


def _digitize_path(path, img):
    """Run the pipeline on an image file and build the editor payload."""
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    r = B.digitize(path)

    fr = r["frame"]
    xcal = r.get("xcal")

    # Build curve data for the frontend
    curves = []
    for c in r["curves"]:
        px = np.asarray(c["px"]).tolist()
        py = np.asarray(c["py"]).tolist()
        role = c.get("role", "emission")

        wn = c.get("wavenumber")
        inten = c.get("intensity")
        wl = (1e7 / wn).tolist() if wn is not None else []
        inten = inten.tolist() if inten is not None else []

        curves.append({
            "role": role,
            "px": px,
            "py": py,
            "wl": wl,
            "intensity": inten,
        })

    # Downscale image for web display (max 1800px wide)
    h, w = img.shape[:2]
    scale = min(1800 / w, 1.0)
    if scale < 1.0:
        img_small = cv2.resize(img, (int(w * scale), int(h * scale)))
    else:
        img_small = img
        scale = 1.0

    _, buf = cv2.imencode(".png", img_small)
    img_b64 = base64.b64encode(buf).decode()

    mol = ""
    try:
        mol = B._ocr_name(gray, fr) or ""
    except Exception:
        pass

    m = B.evaluate(r["cmask"], B.reconstruct_mask(r["cmask"].shape, r["curves"]), fr)

    return {
        "image": img_b64,
        "width": img_small.shape[1],
        "height": img_small.shape[0],
        "scale": scale,
        "frame": {k: int(v * scale) for k, v in fr.items()},
        "frame_orig": {k: int(v) for k, v in fr.items()},
        "curves": curves,
        "molecule": mol,
        "f1": round(m["f1"], 3),
        "strategy": r.get("strategy", ""),
        "has_xcal": xcal is not None,
        "xcal": {
            "a": xcal["a"], "b": xcal["b"],
            "lo": xcal["lo"], "hi": xcal["hi"],
        } if xcal else None,
    }


@app.post("/api/digitize")
async def digitize(file: UploadFile = File(...)):
    raw = await file.read()
    arr = np.frombuffer(raw, np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        return JSONResponse({"error": "Could not decode image"}, status_code=400)
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    tmp.write(raw)
    tmp.close()
    try:
        return _digitize_path(tmp.name, img)
    finally:
        os.unlink(tmp.name)


@app.get("/api/library")
async def library():
    """The bundled Berlman set for slide browsing: [{gid, name}]."""
    entries = []
    names = {}
    order = []
    for cand in SPECTRA_JSON_CANDIDATES:
        try:
            with open(cand) as f:
                raw = json.load(f)
            names = {g: s.get("name", g) for g, s in raw["spectra"].items()}
            order = raw["order"]
            break
        except Exception:
            continue
    if not order and os.path.isdir(PAGES_DIR):
        order = sorted(f[:-4] for f in os.listdir(PAGES_DIR)
                       if f.startswith("graph-") and f.endswith(".png"))
    for gid in order:
        if os.path.exists(os.path.join(PAGES_DIR, gid + ".png")):
            entries.append({"gid": gid, "name": names.get(gid, gid)})
    return {"pages": entries,
            "overlays": os.path.isdir(GALLERY_DIR)}


@app.post("/api/digitize_page")
async def digitize_page(data: dict):
    """Digitize a bundled library page, with a disk cache so revisits are
    instant (first visit runs the full pipeline, ~30-60 s)."""
    gid = re.sub(r"[^A-Za-z0-9\-]", "", data.get("gid") or "")
    path = os.path.join(PAGES_DIR, gid + ".png")
    if not gid or not os.path.exists(path):
        return JSONResponse({"error": f"Unknown page {gid!r}"}, status_code=404)

    cache_path = os.path.join(CACHE, gid + ".json")
    if os.path.exists(cache_path):
        with open(cache_path) as f:
            return json.load(f)

    img = cv2.imread(path, cv2.IMREAD_COLOR)
    payload = _digitize_path(path, img)
    payload["gid"] = gid

    # LRU-ish cache: keep the 40 most recent pages
    with open(cache_path, "w") as f:
        json.dump(payload, f)
    cached = sorted((os.path.join(CACHE, n) for n in os.listdir(CACHE)
                     if n.endswith(".json")), key=os.path.getmtime, reverse=True)
    for old in cached[40:]:
        os.unlink(old)
    return payload


@app.post("/api/export")
async def export(data: dict):
    """Export edited curves as CSV."""
    curves = data.get("curves", [])
    xcal = data.get("xcal")
    frame_orig = data.get("frame_orig", {})
    scale = data.get("scale", 1.0)

    output = io.StringIO()
    w = csv.writer(output)
    w.writerow(["curve", "role", "wavelength_nm", "intensity"])

    for i, c in enumerate(curves):
        role = c.get("role", "emission")
        px_list = c.get("px", [])
        py_list = c.get("py", [])

        yt = frame_orig.get("y_top", 0)
        yb = frame_orig.get("y_bottom", 1)

        for px, py in zip(px_list, py_list):
            real_px = px / scale
            real_py = py / scale
            intensity = (yb - real_py) / (yb - yt) if yb != yt else 0

            if xcal:
                wn = xcal["a"] * real_px + xcal["b"]
                wl = 1e7 / wn if wn > 0 else 0
            else:
                wl = 0

            w.writerow([i, role, f"{wl:.1f}", f"{intensity:.6f}"])

    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode()),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=spectrum.csv"},
    )


def _curve_physical(c, xcal, frame_orig, scale):
    """Edited canvas points -> (wavelength_nm, intensity), sorted by wl."""
    px = np.asarray(c.get("px", []), float) / (scale or 1.0)
    py = np.asarray(c.get("py", []), float) / (scale or 1.0)
    if len(px) < 2:
        return None, None
    yt = frame_orig.get("y_top", 0)
    yb = frame_orig.get("y_bottom", 1)
    inten = (yb - py) / float(yb - yt or 1)
    if not xcal:
        return None, None
    wn = xcal["a"] * px + xcal["b"]
    ok = wn > 0
    if ok.sum() < 2:
        return None, None
    wl = 1e7 / wn[ok]
    inten = inten[ok]
    # Berlman normalization on the RAW points first — a needle apex must be
    # judged before bin-averaging blurs it with its own flank
    mx = float(inten.max()) if len(inten) else 0.0
    if mx >= 0.96:
        inten = inten / mx
    # 0.1 nm dedup-averaging: traced dots are per pixel COLUMN, so a
    # near-vertical needle flank piles many intensities onto one wavelength
    # — exported curves must be single-valued (same rule as the batch
    # pipeline / spectra_all.json).  The bin holding the global peak keeps
    # the apex value instead of the mean, so needles keep their height.
    order = np.argsort(wl)
    wl, inten = wl[order], inten[order]
    bins = np.round(wl, 1)
    uw, idx = np.unique(bins, return_inverse=True)
    sums = np.bincount(idx, weights=inten)
    cnts = np.bincount(idx)
    vals = sums / np.maximum(cnts, 1)
    peak_bin = idx[int(np.argmax(inten))]
    vals[peak_bin] = float(inten.max())
    return uw, vals


def _excel_chart_png(path, wl, inten, title, color):
    """Chart image styled like the Excel scatter charts in the workbooks."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(7.2, 4.2), dpi=130)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(wl, inten, marker="o", ms=2.4, lw=0, color=color)
    ax.set_title(title, fontsize=11, pad=10)
    ax.set_xlabel("Wavelength (nm)", fontsize=9)
    ax.set_ylabel("Normalized Intensity", fontsize=9)
    ax.grid(True, color="#D9D9D9", lw=0.7)
    for s in ax.spines.values():
        s.set_color("#BFBFBF")
    ax.tick_params(colors="#595959", labelsize=8)
    fig.tight_layout()
    fig.savefig(path, facecolor="white")
    plt.close(fig)


def _curve_xlsx(path, wl, inten, title, hexcolor):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.chart import ScatterChart, Reference, Series
    from openpyxl.chart.marker import Marker
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spectrum"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = 'Source: Berlman, "Handbook of Fluorescence Spectra," 2nd Ed., 1971 — digitized by SpectraWolf'
    ws["A2"].font = Font(italic=True, size=9)
    ws.cell(row=4, column=1, value="Wavelength (nm)").font = Font(bold=True)
    ws.cell(row=4, column=2, value="Intensity").font = Font(bold=True)
    for i, (w, v) in enumerate(zip(wl, inten)):
        ws.cell(row=5 + i, column=1, value=round(float(w), 1))
        ws.cell(row=5 + i, column=2, value=round(float(v), 6))
    chart = ScatterChart()
    chart.title = title
    chart.x_axis.title = "Wavelength (nm)"
    chart.y_axis.title = "Normalized Intensity"
    chart.height = 10
    chart.width = 18
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    xref = Reference(ws, min_col=1, min_row=5, max_row=4 + len(wl))
    yref = Reference(ws, min_col=2, min_row=4, max_row=4 + len(wl))
    ser = Series(yref, xref, title_from_data=True)
    ser.marker = Marker(symbol="circle", size=2)
    ser.graphicalProperties.line.noFill = True
    solid = hexcolor.lstrip("#").upper()
    ser.marker.graphicalProperties.solidFill = solid
    ser.marker.graphicalProperties.line.solidFill = solid
    chart.series.append(ser)
    ws.add_chart(chart, "D4")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    wb.save(path)


def _smooth_for_display(wl, inten, n_out=700):
    """Continuous smoothed line for the PhotochemCAD chart.

    Averages duplicate wavelengths, resamples onto a uniform grid, and
    applies a light Savitzky-Golay filter — stroke-width jitter disappears
    while band shapes and needle peaks survive.
    """
    from scipy.signal import savgol_filter
    wl = np.asarray(wl, float)
    inten = np.asarray(inten, float)
    # average duplicates (many pixel columns map to the same 0.1 nm)
    uw, idx = np.unique(np.round(wl, 1), return_inverse=True)
    sums = np.bincount(idx, weights=inten)
    cnts = np.bincount(idx)
    uv = sums / np.maximum(cnts, 1)
    if len(uw) < 8:
        return uw, uv
    grid = np.linspace(uw[0], uw[-1], min(n_out, max(80, len(uw))))
    vals = np.interp(grid, uw, uv)
    win = max(5, min(13, (len(grid) // 40) * 2 + 1))
    try:
        vals = savgol_filter(vals, win, 3)
    except Exception:
        pass
    vals = np.clip(vals, 0, None)
    return grid, vals


def _pccad_html(path, wl, inten, meta, csv_name):
    """Standalone PhotochemCAD-style page: metadata table + hoverable chart."""
    swl, sv = _smooth_for_display(wl, inten)
    pts = [[round(float(w), 2), round(float(v), 5)] for w, v in zip(swl, sv)]
    peak_i = int(np.argmax(inten))
    rows = [
        ("Name", meta["molecule"]),
        ("Spectrum", meta["kind"]),
        ("Peak Wavelength (nm)", f"{wl[peak_i]:.0f}"),
        ("Peak Value", f"{inten[peak_i]:.3f}"),
        ("Solvent", meta.get("solvent", "Cyclohexane (typ.)")),
        ("Reference", "Berlman, 1971 (2nd Ed.)"),
        ("Digitized By", "SpectraWolf"),
        ("Spectrum Data", f'<a class="dl" href="{csv_name}" download>&#11015;&#65039; Download File</a>'),
    ]
    trs = "\n".join(
        f"<tr><th>{k}</th><td>{v}</td></tr>" for k, v in rows)
    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>{meta['molecule']} — {meta['kind']}</title>
<style>
*{{margin:0;padding:0;box-sizing:border-box}}
body{{font-family:-apple-system,'Segoe UI',Arial,sans-serif;background:#fff;color:#212529}}
.topbar{{background:#1e3c96;color:#fff;padding:18px 40px;font-size:22px;font-weight:700;display:flex;align-items:center;gap:12px}}
.topbar .sun{{width:30px;height:30px;border-radius:50%;background:radial-gradient(circle at 50% 50%, #ffd34d 55%, #e8a800 100%);box-shadow:0 0 0 3px #e8a800}}
.navbar{{background:#f4b400;height:14px}}
.page{{max-width:1500px;margin:0 auto;padding:28px 40px}}
h1{{font-weight:400;font-size:34px;margin-bottom:24px}}
h1 .id{{color:#8a8f98}}
h2{{font-weight:400;font-size:26px;margin:26px 0 14px;border-top:1px solid #dee2e6;padding-top:22px}}
.cols{{display:flex;gap:48px;flex-wrap:wrap;align-items:flex-start}}
table{{border-collapse:collapse;min-width:430px}}
th,td{{border:1px solid #dee2e6;padding:11px 16px;text-align:left;font-size:15px}}
th{{font-weight:700;width:230px}}
a.dl{{color:#0d6efd;text-decoration:none;font-weight:600}}
a.dl:hover{{text-decoration:underline}}
.chartbox{{flex:1;min-width:520px;position:relative}}
canvas{{width:100%;height:430px;display:block}}
.tip{{position:absolute;background:#555;color:#fff;padding:6px 10px;border-radius:4px;font-size:13px;pointer-events:none;display:none;white-space:nowrap}}
.footer{{background:#1e3c96;color:#fff;padding:14px 40px;font-size:13px;margin-top:44px;display:flex;justify-content:space-between}}
.footer span{{color:#ffd34d}}
</style></head><body>
<div class="topbar"><div class="sun"></div> SpectraWolf&#8482;</div>
<div class="navbar"></div>
<div class="page">
<h1><span class="id">{meta['gid']}.</span> {meta['molecule']}</h1>
<h2>{meta['kind']} Spectrum</h2>
<div class="cols">
  <table>{trs}</table>
  <div class="chartbox">
    <canvas id="ch"></canvas><div class="tip" id="tip"></div>
  </div>
</div>
</div>
<div class="footer"><div>&#169; SpectraWolf — data digitized from Berlman, <i>Handbook of Fluorescence Spectra</i>, 1971</div><div>Format after <span>PhotochemCAD&#8482;</span> (Lindsey Lab)</div></div>
<script>
const P = {json.dumps(pts)};
const cv = document.getElementById('ch'), tip = document.getElementById('tip');
const C = '#c0392b';
function draw() {{
  const r = cv.getBoundingClientRect();
  cv.width = r.width * devicePixelRatio; cv.height = 430 * devicePixelRatio;
  const g = cv.getContext('2d'); g.scale(devicePixelRatio, devicePixelRatio);
  const W = r.width, H = 430, L = 62, Rt = 14, T = 16, Bm = 46;
  const xs = P.map(p=>p[0]), ys = P.map(p=>p[1]);
  const x0 = Math.min(...xs), x1 = Math.max(...xs);
  const y1 = Math.max(...ys) * 1.06, y0 = 0;
  const X = w => L + (w - x0) / (x1 - x0) * (W - L - Rt);
  const Y = v => T + (1 - (v - y0) / (y1 - y0)) * (H - T - Bm);
  g.clearRect(0,0,W,H);
  g.strokeStyle = '#e3e6ea'; g.lineWidth = 1; g.beginPath();
  const xt = niceTicks(x0, x1, 9), yt = niceTicks(y0, y1, 6);
  g.font = '12px Arial'; g.fillStyle = '#495057';
  for (const t of xt) {{ g.moveTo(X(t), T); g.lineTo(X(t), H - Bm);
    g.textAlign='center'; g.fillText(t, X(t), H - Bm + 18); }}
  for (const t of yt) {{ g.moveTo(L, Y(t)); g.lineTo(W - Rt, Y(t));
    g.textAlign='right'; g.fillText(fmt(t), L - 8, Y(t) + 4); }}
  g.stroke();
  g.strokeStyle = '#adb5bd'; g.strokeRect(L, T, W - L - Rt, H - T - Bm);
  g.strokeStyle = C; g.lineWidth = 2; g.beginPath();
  P.forEach((p,i) => i ? g.lineTo(X(p[0]), Y(p[1])) : g.moveTo(X(p[0]), Y(p[1])));
  g.stroke();
  g.fillStyle = '#495057'; g.textAlign = 'center'; g.font = 'bold 13px Arial';
  g.fillText('Wavelength (nm)', L + (W-L-Rt)/2, H - 8);
  g.save(); g.translate(14, T + (H-T-Bm)/2); g.rotate(-Math.PI/2);
  g.fillText('{meta['ylabel']}', 0, 0); g.restore();
  cv._m = {{X, Y, L, Rt, T, Bm, W, H}};
}}
function niceTicks(a, b, n) {{
  const span = b - a || 1, step0 = span / n, mag = Math.pow(10, Math.floor(Math.log10(step0)));
  const step = [1,2,2.5,5,10].map(m=>m*mag).find(s=>span/s<=n) || mag*10;
  const t = []; for (let v = Math.ceil(a/step)*step; v <= b + 1e-9; v += step) t.push(+v.toFixed(6));
  return t;
}}
function fmt(v) {{ return Math.abs(v) >= 1000 ? v.toLocaleString() : (+v.toFixed(3)).toString(); }}
cv.addEventListener('mousemove', e => {{
  const r = cv.getBoundingClientRect(), m = cv._m; if (!m) return;
  const mx = e.clientX - r.left;
  let best = 0, bd = 1e9;
  for (let i = 0; i < P.length; i++) {{ const d = Math.abs(m.X(P[i][0]) - mx); if (d < bd) {{ bd = d; best = i; }} }}
  const p = P[best];
  tip.style.display = 'block';
  tip.style.left = Math.min(m.X(p[0]) + 12, r.width - 130) + 'px';
  tip.style.top = (m.Y(p[1]) - 40) + 'px';
  tip.innerHTML = p[0].toFixed(2) + '<br><b style="color:#ff8d8d">&#9679;</b> ' + fmt(p[1]);
  draw();
  const g = cv.getContext('2d');
  g.fillStyle = C; g.beginPath(); g.arc(m.X(p[0]), m.Y(p[1]), 5, 0, 7); g.fill();
}});
cv.addEventListener('mouseleave', () => {{ tip.style.display = 'none'; draw(); }});
window.addEventListener('resize', draw);
draw();
</script></body></html>"""
    with open(path, "w") as f:
        f.write(html)


def _safe(s):
    return re.sub(r"[^A-Za-z0-9]+", "_", s or "spectrum").strip("_")[:40] or "spectrum"


@app.post("/api/generate")
async def generate(data: dict):
    """Regenerate all output artifacts from the edited curves.

    Per curve: CSV, Excel workbook (scatter chart), Excel-style chart PNG,
    and a PhotochemCAD-style formatted page.
    """
    curves = data.get("curves", [])
    xcal = data.get("xcal")
    frame_orig = data.get("frame_orig", {})
    scale = data.get("scale", 1.0)
    molecule = data.get("molecule") or "Spectrum"
    gid = data.get("gid") or "SW001"

    if not xcal:
        return JSONResponse({"error": "No x-axis calibration — cannot convert to wavelengths"},
                            status_code=400)

    token = uuid.uuid4().hex[:10]
    outdir = os.path.join(OUTPUTS, token)
    os.makedirs(outdir, exist_ok=True)

    # keep the outputs dir tidy: drop all but the newest 20 sessions
    sessions = sorted((os.path.join(OUTPUTS, d) for d in os.listdir(OUTPUTS)
                       if os.path.isdir(os.path.join(OUTPUTS, d))),
                      key=os.path.getmtime, reverse=True)
    for old in sessions[20:]:
        shutil.rmtree(old, ignore_errors=True)

    results = []
    for c in curves:
        wl, inten = _curve_physical(c, xcal, frame_orig, scale)
        if wl is None:
            continue
        name = c.get("name") or c.get("role") or "curve"
        kind = {"emission": "Emission", "absorption": "Absorption"}.get(
            c.get("role"), name.title())
        color = c.get("color") or {"emission": "#c0392b",
                                   "absorption": "#1e8f4e"}.get(c.get("role"), "#3b82f6")
        base = f"{_safe(molecule)}__{_safe(name)}"

        csv_path = os.path.join(outdir, base + ".csv")
        with open(csv_path, "w", newline="") as fp:
            w = csv.writer(fp)
            w.writerow(["wavelength_nm", "intensity"])
            for wv, iv in zip(wl, inten):
                w.writerow([f"{wv:.1f}", f"{iv:.6f}"])

        title = f"{molecule} — {kind}"
        png_path = os.path.join(outdir, base + "_excel.png")
        _excel_chart_png(png_path, wl, inten, title, color)
        xlsx_path = os.path.join(outdir, base + ".xlsx")
        _curve_xlsx(xlsx_path, wl, inten, title, color)
        pccad_path = os.path.join(outdir, base + "_photochemcad.html")
        _pccad_html(pccad_path, wl, inten,
                    {"molecule": molecule, "kind": kind, "gid": gid,
                     "ylabel": "Molar Extinction (norm.)" if c.get("role") == "absorption"
                               else "Photon Intensity (arb.)"},
                    base + ".csv")

        results.append({
            "name": name,
            "kind": kind,
            "color": color,
            "n_points": int(len(wl)),
            "peak_wl": round(float(wl[int(np.argmax(inten))]), 1),
            "csv": f"/outputs/{token}/{base}.csv",
            "xlsx": f"/outputs/{token}/{base}.xlsx",
            "excel_png": f"/outputs/{token}/{base}_excel.png",
            "pccad": f"/outputs/{token}/{base}_photochemcad.html",
        })

    return {"token": token, "curves": results}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8770)
